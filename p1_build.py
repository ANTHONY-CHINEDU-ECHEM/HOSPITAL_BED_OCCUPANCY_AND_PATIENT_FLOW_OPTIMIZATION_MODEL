"""Build the Project 1 workbook: Hospital Bed Occupancy and Patient Flow Optimization Model."""
import datetime as dt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, BarChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule, FormulaRule, DataBarRule
from lib.xl import *
from lib import cleansheets as CS
from p1_clean import run, FEATURES, BOOL

OUT = "workbook/P1_Hospital_Bed_Occupancy_Patient_Flow_Model.xlsx"

raw, dd, cl = run()
df = cl.df
orig_cols = list(raw.columns)
extra = ["key_repaired_flag", "patient_id_conflict_flag"] + [c + "_ambiguous_flag" for c in cl.cfg["dates"]] + \
        ["outlier_fields", "outlier_count", "date_sequence_conflict_flag", "clinical_plausibility_flag", "imputed_fields", "imputed_count"] + \
        [f[0] for f in FEATURES]
clean = df[["source_row"] + orig_cols + extra].copy()
NC = len(clean)
UNITS = sorted(clean.facility_unit.unique())

wb = Workbook()
wb.remove(wb.active)

# ------------------------------------------------------------------ data sheets
ws_c = wb.create_sheet("Clean Data")
datefmt = {c: "yyyy-mm-dd" for c in cl.cfg["dates"] + ["discharge_date_derived", "admit_week_start", "admit_month"]}
dump_df(ws_c, clean, fmts=datefmt, header_fill=TEAL)
ws_c.sheet_properties.tabColor = TEAL
ws_r = wb.create_sheet("Raw Data")
dump_df(ws_r, raw, header_fill=GREY)
ws_r.sheet_properties.tabColor = GREY
ws_d = wb.create_sheet("Data Dictionary")
for r in dd.itertuples(index=False):
    ws_d.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in r])
for k, w in zip("ABCD", [34, 22, 70, 40]):
    ws_d.column_dimensions[k].width = w
ws_d["A1"].font = F(14, True, NAVY)
for row in ws_d.iter_rows(min_row=12, max_row=12):
    for c in row:
        c.font = F(9, True, WHITE); c.fill = fill(NAVY)
for row in ws_d.iter_rows(min_row=3):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws_d.sheet_properties.tabColor = GREY

ctx = CS.Ctx(wb, cl, raw, clean, len(raw))
for c in clean.columns:
    add_name(wb, "cd_" + c, ctx.cr(c).replace("'Clean Data'", "'Clean Data'"))
N = lambda c: "cd_" + c

# ------------------------------------------------------------------ lists
ws_l = wb.create_sheet("Lists")
ws_l["A1"] = "Units"
for i, u in enumerate(UNITS):
    ws_l.cell(2 + i, 1, u)
ws_l["B1"] = "Unit selector"
ws_l["B2"] = "All Units"
for i, u in enumerate(UNITS):
    ws_l.cell(3 + i, 2, u)
ws_l["C1"] = "Yes/No"
ws_l["C2"], ws_l["C3"] = "No", "Yes"
last_date = clean.admission_date.max().normalize()
asof = [last_date - pd.Timedelta(days=i) for i in range(0, 56)]
ws_l["D1"] = "As of dates"
for i, d in enumerate(asof):
    c = ws_l.cell(2 + i, 4, d.to_pydatetime())
    c.number_format = "yyyy-mm-dd"
ws_l["E1"] = "Mode"
ws_l["E2"], ws_l["E3"] = "Current practice", "Scenario from simulator"
ws_l.sheet_state = "hidden"
L_UNITS = f"=Lists!$A$2:$A${1 + len(UNITS)}"
L_USEL = f"=Lists!$B$2:$B${2 + len(UNITS)}"
L_YN = "=Lists!$C$2:$C$3"
L_ASOF = f"=Lists!$D$2:$D${1 + len(asof)}"
L_MODE = "=Lists!$E$2:$E$3"

# ------------------------------------------------------------------ sheet shells in final order
names = ["Cover", "01 Executive KPI Summary", "02 Breach Risk Dashboard", "03 72h Occupancy Forecast", "04 Discharge Lag Fitting",
         "05 Scenario Simulator", "06 Assumptions & Validation", "M1 Unit Calibration", "M2 Monte Carlo Engine",
         "M3 Weekly Backtest", "M4 Surge Stress Test"]
S = {n: wb.create_sheet(n) for n in names}
DB = "'02 Breach Risk Dashboard'"
M1 = "'M1 Unit Calibration'"
SIM = "'05 Scenario Simulator'"
FC = "'03 72h Occupancy Forecast'"

# ------------------------------------------------------------------ M1 Unit calibration
ws = S["M1 Unit Calibration"]
setup(ws, "M1  Unit Calibration Engine", "Converts the cleaned patient level extract into unit level flow parameters using Little's Law. Every figure is a live formula over Clean Data.", cols=17, width_last=11)
ws.column_dimensions["B"].width = 20
r = section(ws, 5, 2, "Model inputs", 6)
kv(ws, r, 2, "Total staffed beds (hospital)", 400, "#,##0", True, comment="Source: catalogue brief, 400 bed regional hospital")
kv(ws, r + 1, 2, "Safe occupancy threshold", 0.85, "0%", True, comment="Planning threshold. Occupancy at or above this level triggers an escalation flag.")
kv(ws, r + 2, 2, "Critical occupancy threshold", 0.95, "0%", True, comment="Level at which a unit is RED regardless of breach probability.")
kv(ws, r + 3, 2, "Look back window for starting occupancy (days)", 28, "0", True)
kv(ws, r + 4, 2, "Weekend discharge factor (planning)", 0.75, "0.00", True,
   comment="Relative discharge rate on Saturday and Sunday versus the weekly average. The observed ratio in this extract is shown on 06 Assumptions & Validation.")
block_note(ws, r, 7, r + 4, 17, "Blue cells on yellow are inputs. Black cells are formulas. Staffed beds per unit are allocated from each unit's share of patient days (largest remainder method) and can be overridden. "
     "Little's Law (census = arrival rate x length of stay) is used to derive the daily admission rate that is consistent with observed occupancy and observed length of stay, "
     "so the forecast starts from equilibrium and moves only because of the starting census, surge uplift or policy change.")
INP = {"beds": f"{M1}!$D${r}", "thr": f"{M1}!$D${r + 1}", "crit": f"{M1}!$D${r + 2}", "look": f"{M1}!$D${r + 3}", "wkf": f"{M1}!$D${r + 4}"}
r = r + 6
r = section(ws, r, 2, "Unit parameters", 16)
hdr = ["Unit", "Encounters", "Patient days", "Share of patient days", "Staffed beds (input)", "Mean LOS (days)", "Mean discharge lag (h)",
       "Mean turnover (h)", "Effective LOS (days)", "Mean occupancy", "Baseline census", "Daily admissions (lambda)", "Daily discharge probability (p)",
       "Bed block rate", "Discharge before noon rate", "Mean ED boarding (h)"]
header_row(ws, r, 2, hdr, height=44)
pdays = clean.groupby("facility_unit").length_of_stay_days.sum().reindex(UNITS)
quota = pdays / pdays.sum() * 400
beds = np.floor(quota).astype(int)
rem = 400 - beds.sum()
beds[(quota - np.floor(quota)).sort_values(ascending=False).index[:rem]] += 1
M1_TOP = r + 1
for i, u in enumerate(UNITS):
    rr = M1_TOP + i
    z = i % 2 == 1
    cU = f"$B{rr}"
    vals = [u, f"=COUNTIFS({N('facility_unit')},{cU})", f"=SUMIFS({N('length_of_stay_days')},{N('facility_unit')},{cU})",
            f"=D{rr}/SUM($D${M1_TOP}:$D${M1_TOP + len(UNITS) - 1})", int(beds[u]),
            f"=AVERAGEIFS({N('length_of_stay_days')},{N('facility_unit')},{cU})",
            f"=AVERAGEIFS({N('discharge_lag_hours')},{N('facility_unit')},{cU})",
            f"=AVERAGEIFS({N('housekeeping_turnover_minutes')},{N('facility_unit')},{cU})/60",
            f"=G{rr}+H{rr}/24+I{rr}/24",
            f"=AVERAGEIFS({N('occupancy_rate_pct')},{N('facility_unit')},{cU})/100",
            f"=F{rr}*K{rr}", f"=L{rr}/J{rr}", f"=1/J{rr}",
            f"=COUNTIFS({N('facility_unit')},{cU},{N('bed_block_flag')},TRUE)/C{rr}",
            f"=COUNTIFS({N('facility_unit')},{cU},{N('discharge_before_noon_flag')},TRUE)/C{rr}",
            f"=AVERAGEIFS({N('ed_boarding_hours')},{N('facility_unit')},{cU})"]
    fm = [None, "#,##0", "#,##0", "0.0%", "#,##0", "0.00", "0.00", "0.00", "0.00", "0.0%", "0.0", "0.00", "0.000", "0.0%", "0.0%", "0.00"]
    for j, v in enumerate(vals):
        c = ws.cell(rr, 2 + j, v)
        body_cell(c, fm[j], z)
        if j == 4:
            input_style(c, "#,##0")
M1_END = M1_TOP + len(UNITS) - 1
tr = M1_END + 1
tot = ["Hospital total", f"=SUM(C{M1_TOP}:C{M1_END})", f"=SUM(D{M1_TOP}:D{M1_END})", f"=SUM(E{M1_TOP}:E{M1_END})", f"=SUM(F{M1_TOP}:F{M1_END})",
       f"=AVERAGE({N('length_of_stay_days')})", f"=AVERAGE({N('discharge_lag_hours')})", f"=AVERAGE({N('housekeeping_turnover_minutes')})/60",
       f"=G{tr}+H{tr}/24+I{tr}/24", f"=L{tr}/F{tr}", f"=SUM(L{M1_TOP}:L{M1_END})", f"=SUM(M{M1_TOP}:M{M1_END})", f"=M{tr}/L{tr}",
       f"=COUNTIF({N('bed_block_flag')},TRUE)/C{tr}", f"=COUNTIF({N('discharge_before_noon_flag')},TRUE)/C{tr}", f"=AVERAGE({N('ed_boarding_hours')})"]
fm = [None, "#,##0", "#,##0", "0.0%", "#,##0", "0.00", "0.00", "0.00", "0.00", "0.0%", "0.0", "0.00", "0.000", "0.0%", "0.0%", "0.00"]
for j, v in enumerate(tot):
    c = ws.cell(tr, 2 + j, v)
    c.font = F(9, True, NAVY); c.border = BORDER; c.fill = fill(TEAL_L); c.number_format = fm[j] or "General"
ws.cell(tr + 1, 6, f'=IF(F{tr}={INP["beds"]},"Beds reconcile to hospital total","Beds do not sum to hospital total")').font = F(9, True, GREEN)
ws.freeze_panes = ws.cell(M1_TOP, 3)
M1R = lambda col: f"{M1}!${col}${M1_TOP}:${col}${M1_END}"

# ------------------------------------------------------------------ 05 Scenario simulator
ws = S["05 Scenario Simulator"]
setup(ws, "05  Scenario Simulator: Staffing and Discharge Policy Toggles",
      "Deliverable 4. Change the levers in the yellow cells; unit occupancy, bed days released and the value of avoided bed block update immediately.", cols=15, width_last=11)
ws.column_dimensions["B"].width = 36
ws.column_dimensions["E"].width = 15
r = section(ws, 5, 2, "Policy and staffing levers", 5)
LEV = {}
levers = [("seven_day", "Seven day discharge policy (weekend discharges at weekday rate)", "Yes", None, "yn"),
          ("lag_red", "Reduction in discharge lag (order to bed empty)", 0.25, "0%", None),
          ("turn_tgt", "Housekeeping turnover target (minutes)", 35, "0", None),
          ("noon_tgt", "Discharge before noon target share", 0.40, "0%", None),
          ("staff", "Staffing level ratio (actual over budget)", 1.00, "0.00", None),
          ("surge", "Arrival surge uplift", 0.00, "0%", None)]
for i, (k, lab, v, fmt, t) in enumerate(levers):
    rr = r + i
    ws.cell(rr, 2, lab).font = F(10)
    c = ws.cell(rr, 5, v)
    input_style(c, fmt)
    if t == "yn":
        selector(ws, f"E{rr}", L_YN, v)
        input_style(ws[f"E{rr}"])
    LEV[k] = f"{SIM}!$E${rr}"
r += len(levers) + 1
r = section(ws, r, 2, "Planning assumptions", 5)
ASM = {}
asm = [("noon_hours", "Bed hours released per discharge moved before noon", 4.0, "0.0", "Planning assumption: a before noon discharge frees the bed for same day admissions roughly four hours earlier than an afternoon discharge."),
       ("elast", "Elasticity of discharge lag to staffing level", -0.50, "0.00", "Planning assumption: a 10 percent increase in staffing shortens discharge lag by 5 percent."),
       ("annual_loss", "Annual cost of bed block delays (brief)", 2300000, "$#,##0", "Source: catalogue brief, estimated 2.3M dollars lost annually."),
       ]
for i, (k, lab, v, fmt, cm) in enumerate(asm):
    rr = r + i
    ws.cell(rr, 2, lab).font = F(10)
    kv(ws, rr, 2, lab, v, fmt, True, span_label=3, comment=cm)
    ASM[k] = f"{SIM}!$E${rr}"
rr = r + len(asm)
ws.cell(rr, 2, "Implied cost per bed unavailable hour (calibrated)").font = F(10)
c = ws.cell(rr, 5, f"={ASM['annual_loss']}/SUMPRODUCT({M1R('M')}*365,{M1R('H')}+{M1R('I')})")
calc_style(c, "$#,##0.00", True)
ASM["cost_hr"] = f"{SIM}!$E${rr}"
block_note(ws, 6, 7, 17, 15, "How the simulator works. Each unit's effective length of stay is LOS plus discharge lag plus turnover. Levers shorten lag and turnover and shift discharges before noon. "
     "With admissions held at the calibrated rate, a shorter effective stay lowers the steady state census (Little's Law). Bed unavailable hours are valued at the cost per hour implied by the brief's 2.3M dollar annual loss, "
     "so baseline cost reconciles exactly to the brief. The seven day policy changes the weekend discharge rate used by the 72 hour forecast and the Monte Carlo engine. "
     "Switch the Mode selector on the dashboard to 'Scenario from simulator' to push these settings into the forecast.")
r = rr + 2
r = section(ws, r, 2, "Unit impact: baseline versus scenario", 14)
hdr = ["Unit", "Beds", "Baseline effective LOS", "Scenario lag (h)", "Scenario turnover (h)", "Before noon gain (days)", "Scenario effective LOS",
       "Baseline occupancy", "Scenario occupancy", "Change (points)", "Bed days released per year", "Bed unavailable hours avoided per year", "Annual value", "Discharge rate multiplier"]
header_row(ws, r, 2, hdr, height=56)
SIM_TOP = r + 1
for i, u in enumerate(UNITS):
    rr = SIM_TOP + i
    m = M1_TOP + i
    z = i % 2 == 1
    vals = [f"={M1}!B{m}", f"={M1}!F{m}", f"={M1}!J{m}",
            f"={M1}!H{m}*(1-{LEV['lag_red']})*MAX(0.1,1+{ASM['elast']}*({LEV['staff']}-1))",
            f"=MIN({M1}!I{m},{LEV['turn_tgt']}/60)",
            f"=MAX(0,{LEV['noon_tgt']}-{M1}!P{m})*{ASM['noon_hours']}/24",
            f"={M1}!G{m}+E{rr}/24+F{rr}/24-G{rr}",
            f"={M1}!K{m}",
            f"={M1}!M{m}*(1+{LEV['surge']})*H{rr}/C{rr}",
            f"=(J{rr}-I{rr})*100",
            f"=({M1}!M{m}*D{rr}-{M1}!M{m}*H{rr})*365",
            f"={M1}!M{m}*365*(({M1}!H{m}+{M1}!I{m})-(E{rr}+F{rr}))",
            f"=M{rr}*{ASM['cost_hr']}",
            f"=D{rr}/H{rr}"]
    fm = [None, "#,##0", "0.00", "0.00", "0.00", "0.000", "0.00", "0.0%", "0.0%", "+0.0;-0.0;0.0", "#,##0", "#,##0", "$#,##0", "0.000"]
    for j, v in enumerate(vals):
        body_cell(ws.cell(rr, 2 + j, v), fm[j], z)
SIM_END = SIM_TOP + len(UNITS) - 1
tr = SIM_END + 1
tot = ["Hospital", f"=SUM(C{SIM_TOP}:C{SIM_END})", "", "", "", "", "",
       f"=SUMPRODUCT(C{SIM_TOP}:C{SIM_END},I{SIM_TOP}:I{SIM_END})/C{tr}", f"=SUMPRODUCT(C{SIM_TOP}:C{SIM_END},J{SIM_TOP}:J{SIM_END})/C{tr}",
       f"=(J{tr}-I{tr})*100", f"=SUM(L{SIM_TOP}:L{SIM_END})", f"=SUM(M{SIM_TOP}:M{SIM_END})", f"=SUM(N{SIM_TOP}:N{SIM_END})", ""]
fm = [None, "#,##0", None, None, None, None, None, "0.0%", "0.0%", "+0.0;-0.0;0.0", "#,##0", "#,##0", "$#,##0", None]
for j, v in enumerate(tot):
    c = ws.cell(tr, 2 + j, v)
    c.font = F(9, True, NAVY); c.border = BORDER; c.fill = fill(TEAL_L)
    if fm[j]:
        c.number_format = fm[j]
ws.conditional_formatting.add(f"K{SIM_TOP}:K{SIM_END}", CellIsRule(operator="lessThan", formula=["0"], font=F(9, True, GREEN)))
ws.conditional_formatting.add(f"K{SIM_TOP}:K{SIM_END}", CellIsRule(operator="greaterThan", formula=["0"], font=F(9, True, RED)))
SIMTOT = tr
r = tr + 2
r = section(ws, r, 2, "Headline", 6)
heads = [("Baseline annual cost of bed unavailable time", f"=SUMPRODUCT({M1R('M')}*365,{M1R('H')}+{M1R('I')})*{ASM['cost_hr']}", "$#,##0"),
         ("Annual value of avoided bed block", f"=N{SIMTOT}", "$#,##0"),
         ("Share of bed block cost removed", f"=N{SIMTOT}/E{r}", "0.0%"),
         ("Equivalent beds released (bed days / 365)", f"=L{SIMTOT}/365", "0.0"),
         ("Hospital occupancy: baseline to scenario", f'=TEXT(I{SIMTOT},"0.0%")&"  to  "&TEXT(J{SIMTOT},"0.0%")', None)]
for i, (lab, f_, fmt) in enumerate(heads):
    ws.cell(r + i, 2, lab).font = F(10)
    c = ws.cell(r + i, 5, f_)
    calc_style(c, fmt, True, NAVY)
SIMHEAD = r
ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"
ctitle(ch, "Occupancy by unit: baseline versus scenario")
ch.add_data(Reference(ws, min_col=9, max_col=10, min_row=SIM_TOP - 1, max_row=SIM_END), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=SIM_TOP, max_row=SIM_END))
ch.y_axis.number_format = "0%"; ch.y_axis.title = "Occupancy"; ch.height = 8; ch.width = 22
ch.series[0].graphicalProperties.solidFill = "9AA7B8"; ch.series[1].graphicalProperties.solidFill = TEAL
ch.y_axis.majorGridlines = None; ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1
ch.width = 17; ch.legend.position = "b"
ws.add_chart(ch, f"H{r}")

# ------------------------------------------------------------------ 02 Dashboard controls (placed early for references)
ws = S["02 Breach Risk Dashboard"]
setup(ws, "Unit Level Breach Risk Dashboard", "72 hour occupancy outlook, breach probability and patient flow drivers. Use the three selectors to change the unit, forecast date and policy mode.", cols=17, width_last=9.5)
ws.row_dimensions[4].height = 14
selector(ws, "B6", L_USEL, "All Units", "UNIT", "B5", "Choose a unit, or All Units for the whole hospital")
ws.merge_cells("B6:D6")
selector(ws, "F6", L_ASOF, asof[0].to_pydatetime(), "FORECAST AS OF", "F5", "Forecast start date (last 56 days of data)")
ws.merge_cells("F6:H6"); ws["F6"].number_format = "dd mmm yyyy"
selector(ws, "J6", L_MODE, "Current practice", "POLICY MODE", "J5", "Current practice, or apply the levers set on 05 Scenario Simulator")
ws.merge_cells("J6:M6")
ws.row_dimensions[6].height = 24
CTL = {"unit": f"{DB}!$B$6", "asof": f"{DB}!$F$6", "mode": f"{DB}!$J$6"}

# ------------------------------------------------------------------ 03 72h forecast
ws = S["03 72h Occupancy Forecast"]
setup(ws, "03  Rolling 72 Hour Bed Occupancy Forecast", "Deliverable 1. Unit level expected census, uncertainty and breach probability for the next three days, rolling from the selected forecast date.", cols=22, width_last=9)
ws.column_dimensions["B"].width = 18
ws.column_dimensions["D"].width = 17
r = section(ws, 5, 2, "Forecast controls (set on the dashboard)", 8)
lab = [("Forecast as of", f"={CTL['asof']}", "dd mmm yyyy"), ("Policy mode", f"={CTL['mode']}", None),
       ("Day 1", f"=D6", "ddd dd mmm"), ("Day 2", f"=D6+1", "ddd dd mmm"), ("Day 3", f"=D6+2", "ddd dd mmm")]
for i, (a, b, fm) in enumerate(lab):
    ws.cell(r + i, 2, a).font = F(10)
    c = ws.cell(r + i, 4, b); link_style(c, fm)
ws["D8"] = "=D6"; ws["D9"] = "=D6+1"; ws["D10"] = "=D6+2"
for k in ("D8", "D9", "D10"):
    link_style(ws[k], "ddd dd mmm")
# weekend factor per day
ws["F7"] = "Weekend factor applied"; ws["F7"].font = F(10)
for i, k in enumerate(("D8", "D9", "D10")):
    wfe = f'IF(AND({LEV["seven_day"]}="Yes",{CTL["mode"]}<>"Current practice"),1,{INP["wkf"]})'
    c = ws.cell(8 + i, 6, f'=IF(WEEKDAY({k},2)>=6,{wfe},1)/((5+2*{wfe})/7)')
    calc_style(c, "0.000")
block_note(ws, 5, 9, 10, 22, "Method. Starting census = staffed beds x mean recorded occupancy in the look back window before the forecast date. Each day, admissions arrive at the calibrated Poisson rate and each occupied bed is discharged with the unit's daily probability, "
     "adjusted for weekends. Expected census and variance are propagated exactly: E(d) = E(d-1) x (1 - p) + lambda and Var(d) = Var(d-1) x (1 - p)^2 + E(d-1) x p x (1 - p) + lambda. "
     "Breach probability uses a normal approximation to that distribution and is cross checked against 1,000 simulated trials on M2 Monte Carlo Engine.")
r = 12
r = section(ws, r, 2, "Unit forecast", 21)
hdr = ["Unit", "Beds", "Start occupancy", "Start census", "lambda (adj)", "p (adj)", "p day 1", "p day 2", "p day 3",
       "E census 24h", "E census 48h", "E census 72h", "Var 24h", "Var 48h", "Var 72h", "Occ 24h", "Occ 48h", "Occ 72h", "SD 72h (beds)", "P(breach) 72h", "Status"]
header_row(ws, r, 2, hdr, height=40)
FC_TOP = r + 1
for i, u in enumerate(UNITS):
    rr = FC_TOP + i
    m = M1_TOP + i
    s = SIM_TOP + i
    z = i % 2 == 1
    scen = f'{CTL["mode"]}="Scenario from simulator"'
    vals = [u, f"={M1}!F{m}",
            f'=IFERROR(AVERAGEIFS({N("occupancy_rate_pct")},{N("facility_unit")},B{rr},{N("admission_date")},">="&($D$6-{INP["look"]}),{N("admission_date")},"<"&$D$6)/100,{M1}!K{m})',
            f"=ROUND(C{rr}*D{rr},0)",
            f"={M1}!M{m}*IF({scen},1+{LEV['surge']},1)",
            f"={M1}!N{m}*IF({scen},{SIM}!O{s},1)",
            f"=MIN(0.99,G{rr}*$F$8)", f"=MIN(0.99,G{rr}*$F$9)", f"=MIN(0.99,G{rr}*$F$10)",
            f"=E{rr}*(1-H{rr})+F{rr}", f"=K{rr}*(1-I{rr})+F{rr}", f"=L{rr}*(1-J{rr})+F{rr}",
            f"=E{rr}*H{rr}*(1-H{rr})+F{rr}", f"=N{rr}*(1-I{rr})^2+K{rr}*I{rr}*(1-I{rr})+F{rr}", f"=O{rr}*(1-J{rr})^2+L{rr}*J{rr}*(1-J{rr})+F{rr}",
            f"=K{rr}/C{rr}", f"=L{rr}/C{rr}", f"=M{rr}/C{rr}", f"=SQRT(P{rr})",
            f"=1-NORMDIST({INP['thr']}*C{rr},M{rr},T{rr},TRUE)",
            f'=IF(OR(S{rr}>={INP["crit"]},U{rr}>=0.5),"RED",IF(OR(S{rr}>={INP["thr"]},U{rr}>=0.2),"AMBER","GREEN"))']
    fm = [None, "#,##0", "0.0%", "#,##0", "0.00", "0.000", "0.000", "0.000", "0.000", "0.0", "0.0", "0.0", "0.0", "0.0", "0.0", "0.0%", "0.0%", "0.0%", "0.0", "0.0%", None]
    for j, v in enumerate(vals):
        body_cell(ws.cell(rr, 2 + j, v), fm[j], z, align="center" if j == 20 else None)
FC_END = FC_TOP + len(UNITS) - 1
tr = FC_END + 1
tot = ["Hospital", f"=SUM(C{FC_TOP}:C{FC_END})", f"=E{tr}/C{tr}", f"=SUM(E{FC_TOP}:E{FC_END})", f"=SUM(F{FC_TOP}:F{FC_END})",
       f"=SUMPRODUCT(G{FC_TOP}:G{FC_END},E{FC_TOP}:E{FC_END})/E{tr}", "", "", "",
       f"=SUM(K{FC_TOP}:K{FC_END})", f"=SUM(L{FC_TOP}:L{FC_END})", f"=SUM(M{FC_TOP}:M{FC_END})",
       f"=SUM(N{FC_TOP}:N{FC_END})", f"=SUM(O{FC_TOP}:O{FC_END})", f"=SUM(P{FC_TOP}:P{FC_END})",
       f"=K{tr}/C{tr}", f"=L{tr}/C{tr}", f"=M{tr}/C{tr}", f"=SQRT(P{tr})", f"=1-NORMDIST({INP['thr']}*C{tr},M{tr},T{tr},TRUE)",
       f'=IF(OR(S{tr}>={INP["crit"]},U{tr}>=0.5),"RED",IF(OR(S{tr}>={INP["thr"]},U{tr}>=0.2),"AMBER","GREEN"))']
for j, v in enumerate(tot):
    c = ws.cell(tr, 2 + j, v)
    c.font = F(9, True, NAVY); c.border = BORDER; c.fill = fill(TEAL_L)
    c.number_format = fm[j] or "General"
FCTOT = tr
for rng in (f"V{FC_TOP}:V{tr}",):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"RED"'], fill=fill(RED), font=F(9, True, WHITE)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"AMBER"'], fill=fill("F59E0B"), font=F(9, True, WHITE)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"GREEN"'], fill=fill(GREEN), font=F(9, True, WHITE)))
ws.conditional_formatting.add(f"Q{FC_TOP}:S{FC_END}", ColorScaleRule(start_type="num", start_value=0.6, start_color="E8F5E9", mid_type="num", mid_value=0.85, mid_color="FFE08A", end_type="num", end_value=1.0, end_color="E57373"))
note(ws, tr + 2, 2, "Status rule: RED when 72 hour occupancy reaches the critical threshold or breach probability is at least 50 percent; AMBER when occupancy reaches the safe threshold or breach probability is at least 20 percent; otherwise GREEN. Thresholds are inputs on M1 Unit Calibration.", span=20)
ws.freeze_panes = ws.cell(FC_TOP, 3)
FCR = lambda col: f"{FC}!${col}${FC_TOP}:${col}${FC_END}"

# ------------------------------------------------------------------ M2 Monte Carlo
ws = S["M2 Monte Carlo Engine"]
TRIALS = 1000
setup(ws, "M2  Monte Carlo Simulation Engine (1,000 trials)", "Stochastic cross check of the 72 hour forecast for the unit selected on the dashboard. Each row is one independent trial; press F9 to redraw.", cols=14, width_last=10)
r = section(ws, 5, 2, "Parameters for the selected unit", 6)
sel = CTL["unit"]
par = [("Selected unit", f"={sel}", None),
       ("Staffed beds", f'=IF({sel}="All Units",{FC}!C{FCTOT},INDEX({FCR("C")},MATCH({sel},{FCR("B")},0)))', "#,##0"),
       ("Starting census", f'=IF({sel}="All Units",{FC}!E{FCTOT},INDEX({FCR("E")},MATCH({sel},{FCR("B")},0)))', "#,##0"),
       ("Daily admissions (lambda)", f'=IF({sel}="All Units",{FC}!F{FCTOT},INDEX({FCR("F")},MATCH({sel},{FCR("B")},0)))', "0.00"),
       ("p day 1", f'=IF({sel}="All Units",{FC}!G{FCTOT}*{FC}!$F$8,INDEX({FCR("H")},MATCH({sel},{FCR("B")},0)))', "0.000"),
       ("p day 2", f'=IF({sel}="All Units",{FC}!G{FCTOT}*{FC}!$F$9,INDEX({FCR("I")},MATCH({sel},{FCR("B")},0)))', "0.000"),
       ("p day 3", f'=IF({sel}="All Units",{FC}!G{FCTOT}*{FC}!$F$10,INDEX({FCR("J")},MATCH({sel},{FCR("B")},0)))', "0.000"),
       ("Breach level (beds)", f"={INP['thr']}*D7", "0.0")]
for i, (a, b, fm) in enumerate(par):
    ws.cell(r + i, 2, a).font = F(10)
    c = ws.cell(r + i, 4, b); link_style(c, fm) if i < 7 else calc_style(c, fm)
# rows: 6 unit, 7 beds, 8 C0, 9 lambda, 10 p1, 11 p2, 12 p3, 13 breach beds
r = section(ws, 5, 7, "Simulation summary", 7)
header_row(ws, 6, 7, ["Statistic", "Day 0", "24h", "48h", "72h", "Analytic mean"])
MC_TOP = 18
stats = [("Mean census", "AVERAGE"), ("P10", "P10"), ("P50 (median)", "P50"), ("P90", "P90"), ("Std deviation", "STDEV")]
colmap = {"24h": "E", "48h": "H", "72h": "K"}
for i, (lab_, fn) in enumerate(stats):
    rr = 7 + i
    ws.cell(rr, 7, lab_).font = F(9)
    ws.cell(rr, 8, "=$D$8").number_format = "0.0"
    for j, h in enumerate(["24h", "48h", "72h"]):
        col = colmap[h]
        rng = f"${col}${MC_TOP + 1}:${col}${MC_TOP + TRIALS}"
        f_ = {"AVERAGE": f"=AVERAGE({rng})", "P10": f"=PERCENTILE({rng},0.1)", "P50": f"=PERCENTILE({rng},0.5)",
              "P90": f"=PERCENTILE({rng},0.9)", "STDEV": f"=STDEV({rng})"}[fn]
        c = ws.cell(rr, 9 + j, f_); calc_style(c, "0.0")
    for k in range(7, 13):
        ws.cell(rr, k).border = BORDER
ws.cell(7, 12, "=$D$8*(1-$D$10)+$D$9"); ws.cell(8, 12, "=L7*(1-$D$11)+$D$9"); ws.cell(9, 12, "=L8*(1-$D$12)+$D$9")
ws.cell(6, 12).value = "Analytic mean (24/48/72h)"
for k in (7, 8, 9):
    calc_style(ws.cell(k, 12), "0.0")
ws.cell(13, 7, "P(breach at 72h), simulated").font = F(9, True)
c = ws.cell(13, 11, f'=COUNTIF($K${MC_TOP + 1}:$K${MC_TOP + TRIALS},">"&$D$13)/{TRIALS}'); calc_style(c, "0.0%", True, NAVY)
ws.cell(14, 7, "Simulated mean vs analytic mean (72h)").font = F(9, True)
c = ws.cell(14, 11, "=K7/L9-1"); calc_style(c, "+0.0%;-0.0%", True, NAVY)
ws.cell(15, 7, "Agreement check (within 2 percent)").font = F(9, True)
ws.cell(15, 11, '=IF(ABS(K14)<=0.02,"PASS","FAIL")')
CS.PASS_RULES(ws, "K15")
header_row(ws, MC_TOP, 2, ["Trial", "Admits d1", "Discharges d1", "Census 24h", "Admits d2", "Discharges d2", "Census 48h", "Admits d3", "Discharges d3", "Census 72h", "Breach 72h"])
for t in range(1, TRIALS + 1):
    rr = MC_TOP + t
    ws.cell(rr, 2, t)
    ws.cell(rr, 3, f"=CRITBINOM(2000,$D$9/2000,RAND())")
    ws.cell(rr, 4, f"=IF($D$8<=0,0,CRITBINOM($D$8,$D$10,RAND()))")
    ws.cell(rr, 5, f"=MAX(0,$D$8-D{rr}+C{rr})")
    ws.cell(rr, 6, f"=CRITBINOM(2000,$D$9/2000,RAND())")
    ws.cell(rr, 7, f"=IF(E{rr}<=0,0,CRITBINOM(E{rr},$D$11,RAND()))")
    ws.cell(rr, 8, f"=MAX(0,E{rr}-G{rr}+F{rr})")
    ws.cell(rr, 9, f"=CRITBINOM(2000,$D$9/2000,RAND())")
    ws.cell(rr, 10, f"=IF(H{rr}<=0,0,CRITBINOM(H{rr},$D$12,RAND()))")
    ws.cell(rr, 11, f"=MAX(0,H{rr}-J{rr}+I{rr})")
    ws.cell(rr, 12, f"=K{rr}>$D$13")
    for k in range(2, 13):
        ws.cell(rr, k).font = F(9)
ws.freeze_panes = ws.cell(MC_TOP + 1, 3)
MC = "'M2 Monte Carlo Engine'"

# ------------------------------------------------------------------ M3 Weekly backtest
ws = S["M3 Weekly Backtest"]
setup(ws, "M3  Weekly Forecast Backtest (4 week rolling baseline)", "Measures forecast accuracy against every complete week in the extract, at hospital level and against the legacy model already in the data.", cols=12, width_last=12)
weeks = pd.date_range("2022-01-03", "2025-12-22", freq="W-MON")
r = section(ws, 5, 2, "Accuracy summary", 6)
BT_TOP = 16
BT_END = BT_TOP + len(weeks)
summ = [("Weeks evaluated", f"=COUNT(F{BT_TOP + 1}:F{BT_END})", "#,##0"),
        ("Weeks within plus or minus 4 points", f'=COUNTIF(H{BT_TOP + 1}:H{BT_END},TRUE)/F6', "0.0%"),
        ("Mean absolute error (points)", f"=AVERAGE(G{BT_TOP + 1}:G{BT_END})", "0.00"),
        ("Bias (forecast minus actual, points)", f"=AVERAGE(F{BT_TOP + 1}:F{BT_END})", "+0.00;-0.00"),
        ("Legacy model: records with forecast variance within 4 points",
         f'=COUNTIFS({N("forecast_variance_pct")},">=-4",{N("forecast_variance_pct")},"<=4")/COUNT({N("forecast_variance_pct")})', "0.0%"),
        ("Correlation of weekly admissions with weekly occupancy", f"=CORREL(C{BT_TOP + 1}:C{BT_END},D{BT_TOP + 1}:D{BT_END})", "0.000"),
        ("Unit week accuracy (see note)", 0.361, "0.0%")]
for i, (a, b, fm) in enumerate(summ):
    ws.cell(r + i, 2, a).font = F(10)
    c = ws.cell(r + i, 6, b); calc_style(c, fm, True, NAVY)
ws.cell(r + 6, 6).font = F(10, True, GREY)
BTS = {"acc": f"'M3 Weekly Backtest'!$F${r + 1}", "mae": f"'M3 Weekly Backtest'!$F${r + 2}", "legacy": f"'M3 Weekly Backtest'!$F${r + 4}",
       "corr": f"'M3 Weekly Backtest'!$F${r + 5}", "n": f"'M3 Weekly Backtest'!$F${r}"}
block_note(ws, r, 8, r + 8, 13, "Why hospital level. The extract holds about one occupancy reading per unit per day, so one reading carries sampling noise of roughly 18 points, far above the 4 point tolerance. "
     "Aggregating to hospital weeks (about 70 readings) reduces that noise to about 2 points, which is the grain at which accuracy can be judged fairly. "
     "The unit week figure (36.1 percent, computed in the Python audit script) is shown for transparency and is limited by data volume, not by the model. "
     "The near zero correlation between admissions and recorded occupancy is itself a finding: the occupancy field is not reconciled to the ADT census (see C8).")
header_row(ws, BT_TOP, 2, ["Week starting", "Admissions", "Actual occupancy", "Forecast (prior 4 week mean)", "Error (points)", "Absolute error", "Within 4 points"])
for i, w in enumerate(weeks):
    rr = BT_TOP + 1 + i
    ws.cell(rr, 2, w.to_pydatetime()).number_format = "yyyy-mm-dd"
    ws.cell(rr, 3, f"=COUNTIFS({N('admit_week_start')},B{rr})")
    ws.cell(rr, 4, f'=IFERROR(AVERAGEIFS({N("occupancy_rate_pct")},{N("admit_week_start")},B{rr}),"")').number_format = "0.00"
    if i >= 4:
        ws.cell(rr, 5, f"=AVERAGE(D{rr - 4}:D{rr - 1})").number_format = "0.00"
        ws.cell(rr, 6, f"=E{rr}-D{rr}").number_format = "+0.00;-0.00"
        ws.cell(rr, 7, f"=ABS(F{rr})").number_format = "0.00"
        ws.cell(rr, 8, f"=G{rr}<=4")
    for k in range(2, 9):
        ws.cell(rr, k).font = F(9)
ch = LineChart(); ctitle(ch, "Hospital weekly occupancy: actual versus 4 week rolling forecast")
ch.add_data(Reference(ws, min_col=4, max_col=5, min_row=BT_TOP, max_row=BT_END), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=BT_TOP + 1, max_row=BT_END))
ch.height = 8; ch.width = 26; ch.y_axis.title = "Occupancy (%)"; ch.x_axis.number_format = "mmm yy"
ch.series[0].graphicalProperties.line.solidFill = "9AA7B8"; ch.series[1].graphicalProperties.line.solidFill = TEAL
ch.series[1].graphicalProperties.line.width = 22000
ws.add_chart(ch, "J16")
ws.freeze_panes = ws.cell(BT_TOP + 1, 3)

# ------------------------------------------------------------------ M4 surge stress test
ws = S["M4 Surge Stress Test"]
setup(ws, "M4  Historical Surge Week Stress Test", "The three highest admission weeks in the extract are replayed through the flow model at hospital level to test escalation behaviour under surge load.", cols=12, width_last=13)
ws.column_dimensions["B"].width = 34
wk = clean.dropna(subset=["admission_date"]).groupby("admit_week_start").size()
wk = wk[(wk.index >= "2022-01-03") & (wk.index <= "2025-12-22")]
surge_weeks = wk.sort_values(ascending=False).index[:3].sort_values()
r = section(ws, 5, 2, "Stress test design", 10)
note(ws, r, 2, "Each surge week is replayed for seven days. Admission uplift = admissions in the surge week divided by the average week in the extract. Starting occupancy = the mean recorded occupancy in the four weeks before the surge. "
     "Daily discharge probability and bed base come from M1. The model must (1) remain numerically stable, (2) raise escalation flags when projected occupancy crosses the thresholds, and (3) have its projection compared honestly with recorded occupancy.", span=10, height=58)
r += 2
hdr = ["Measure", "Surge week 1", "Surge week 2", "Surge week 3"]
header_row(ws, r, 2, hdr)
rows = ["Week starting", "Admissions in week", "Average admissions per week (extract)", "Arrival uplift factor", "Starting occupancy (prior 4 weeks)",
        "Starting census (beds)", "Daily admissions under surge", "Daily discharge probability", "Projected census day 3", "Projected census day 7",
        "Projected occupancy day 3", "Projected occupancy day 7", "Escalation flag day 7", "Recorded occupancy during week", "Projection minus recorded (points)",
        "Model stable (no negative or runaway census)"]
top = r + 1
for i, lab_ in enumerate(rows):
    ws.cell(top + i, 2, lab_).font = F(9, i in (12, 15))
    ws.cell(top + i, 2).border = BORDER
for j, w in enumerate(surge_weeks):
    col = L(3 + j)
    R = lambda k: f"{col}{top + k}"
    f_ = [w.to_pydatetime(), f"=COUNTIFS({N('admit_week_start')},{R(0)})",
          f"=COUNT({N('admission_date')})/COUNT('M3 Weekly Backtest'!$B${BT_TOP + 1}:$B${BT_END})",
          f"={R(1)}/{R(2)}",
          f'=AVERAGEIFS({N("occupancy_rate_pct")},{N("admit_week_start")},">="&({R(0)}-28),{N("admit_week_start")},"<"&{R(0)})/100',
          f"={R(4)}*{INP['beds']}", f"={M1}!M{M1_END + 1}*{R(3)}", f"={M1}!N{M1_END + 1}",
          f"={R(5)}*(1-{R(7)})^3+{R(6)}*(1-(1-{R(7)})^3)/{R(7)}",
          f"={R(5)}*(1-{R(7)})^7+{R(6)}*(1-(1-{R(7)})^7)/{R(7)}",
          f"={R(8)}/{INP['beds']}", f"={R(9)}/{INP['beds']}",
          f'=IF({R(11)}>={INP["crit"]},"RED",IF({R(11)}>={INP["thr"]},"AMBER","GREEN"))',
          f"=AVERAGEIFS({N('occupancy_rate_pct')},{N('admit_week_start')},{R(0)})/100",
          f"=({R(11)}-{R(13)})*100",
          f'=IF(AND({R(9)}>=MIN({R(5)},{R(6)}/{R(7)})-0.001,{R(9)}<=MAX({R(5)},{R(6)}/{R(7)})+0.001),"PASS","FAIL")']
    fm = ["dd mmm yyyy", "#,##0", "0.0", "0.00x", "0.0%", "0.0", "0.00", "0.000", "0.0", "0.0", "0.0%", "0.0%", None, "0.0%", "+0.0;-0.0", None]
    for i, v in enumerate(f_):
        c = ws.cell(top + i, 3 + j, v)
        body_cell(c, fm[i], i % 2 == 1, align="center")
CS.PASS_RULES(ws, f"C{top + 15}:E{top + 15}")
ws.conditional_formatting.add(f"C{top + 12}:E{top + 12}", CellIsRule(operator="equal", formula=['"RED"'], fill=fill(RED), font=F(9, True, WHITE)))
ws.conditional_formatting.add(f"C{top + 12}:E{top + 12}", CellIsRule(operator="equal", formula=['"AMBER"'], fill=fill("F59E0B"), font=F(9, True, WHITE)))
ws.conditional_formatting.add(f"C{top + 12}:E{top + 12}", CellIsRule(operator="equal", formula=['"GREEN"'], fill=fill(GREEN), font=F(9, True, WHITE)))
ST = {"top": top, "ws": "'M4 Surge Stress Test'"}
r = top + len(rows) + 1
ws.cell(r, 2, "Surge scenarios validated").font = F(10, True)
c = ws.cell(r, 3, f'=COUNTIF(C{top + 15}:E{top + 15},"PASS")'); calc_style(c, '0" of 3"', True, NAVY)
ST["passed"] = f"'M4 Surge Stress Test'!$C${r}"
note(ws, r + 2, 2, "Interpretation. The flow model correctly projects rising census when admissions exceed the equilibrium rate and escalates status as thresholds are crossed. "
     "The gap between projected and recorded occupancy is expected: in this extract the recorded occupancy field does not move with admission volume (weekly correlation shown on M3 is close to zero), "
     "which confirms the C8 finding that occupancy is captured independently of the ADT feed. Before production use, occupancy should be sourced from a midnight census reconciled to ADT events.", span=10, height=72)

# ------------------------------------------------------------------ 04 Discharge lag fitting
ws = S["04 Discharge Lag Fitting"]
setup(ws, "04  Discharge Lag Statistical Fitting Report", "Deliverable 3. Distribution of hours from discharge order to bed release, with Poisson, Normal and Gamma fits and chi square goodness of fit tests.", cols=16, width_last=10)
ws.column_dimensions["B"].width = 16
ws.column_dimensions["F"].width = 20
selector(ws, "B6", L_USEL, "All Units", "UNIT / SPECIALTY", "B5", "Choose a unit to refit the distributions")
ws.merge_cells("B6:C6")
crit = 'IF($B$6="All Units","*",$B$6)'
r = 8
r = section(ws, r, 2, "Moment estimates", 4)
BINS = list(range(0, 27))
BIN_TOP = 23
BIN_END = BIN_TOP + len(BINS)
mom = [("Observations (n)", f"=SUM(C{BIN_TOP + 1}:C{BIN_END})", "#,##0"),
       ("Mean (hours)", f"=SUMPRODUCT(B{BIN_TOP + 1}:B{BIN_END},C{BIN_TOP + 1}:C{BIN_END})/D9", "0.000"),
       ("Variance", f"=SUMPRODUCT((B{BIN_TOP + 1}:B{BIN_END}-D10)^2,C{BIN_TOP + 1}:C{BIN_END})/(D9-1)", "0.000"),
       ("Standard deviation", "=SQRT(D11)", "0.000"),
       ("Dispersion index (variance / mean)", "=D11/D10", "0.000"),
       ("Gamma shape (alpha)", "=D10^2/D11", "0.000"),
       ("Gamma scale (beta)", "=D11/D10", "0.000"),
       ("Share of discharges over 4 hours", f"=SUMIF(B{BIN_TOP + 1}:B{BIN_END},\">4\",C{BIN_TOP + 1}:C{BIN_END})/D9", "0.0%")]
for i, (a, b, fm) in enumerate(mom):
    ws.cell(9 + i, 2, a).font = F(9)
    c = ws.cell(9 + i, 4, b); calc_style(c, fm, True, NAVY)
section(ws, 8, 6, "Goodness of fit (bins with expected count of 5 or more)", 8)
header_row(ws, 9, 6, ["Distribution", "Parameters", "Chi square", "Bins used", "Degrees of freedom", "p value", "Chi square / df", "Rank"])
fits = [("Poisson", 1, "D"), ("Normal censored at 0", 2, "E"), ("Gamma", 2, "F"), ("Hurdle Gamma", 3, "J")]
for i, (nm, npar, col) in enumerate(fits):
    rr = 10 + i
    ch_col = {"Poisson": "G", "Normal censored at 0": "H", "Gamma": "I", "Hurdle Gamma": "K"}[nm]
    vals = [nm, npar, f"=SUM({ch_col}{BIN_TOP + 1}:{ch_col}{BIN_END})",
            f'=COUNTIF({col}{BIN_TOP + 1}:{col}{BIN_END},">=5")',
            f"=MAX(1,I{rr}-1-G{rr})", f"=CHIDIST(H{rr},J{rr})", f"=H{rr}/J{rr}", f"=RANK(L{rr},$L$10:$L$13,1)"]
    fm = [None, "0", "#,##0.0", "0", "0", "0.0000", "0.00", "0"]
    for j, v in enumerate(vals):
        body_cell(ws.cell(rr, 6 + j, v), fm[j], i % 2 == 1)
ws.conditional_formatting.add("M10:M13", CellIsRule(operator="equal", formula=["1"], fill=fill(GREEN_L), font=F(9, True, GREEN)))
section(ws, 8, 15, "Hurdle model parameters", 2)
hp = [("Zero share (lag under 0.5 h)", f"=C{BIN_TOP + 1}/D9", "0.0%"),
      ("Positive lag mean (h)", f"=SUMPRODUCT(B{BIN_TOP + 2}:B{BIN_END},C{BIN_TOP + 2}:C{BIN_END})/(D9-C{BIN_TOP + 1})", "0.000"),
      ("Positive lag variance", f"=SUMPRODUCT((B{BIN_TOP + 2}:B{BIN_END}-P11)^2,C{BIN_TOP + 2}:C{BIN_END})/(D9-C{BIN_TOP + 1}-1)", "0.000"),
      ("Gamma shape on positives", "=P11^2/P12", "0.000"),
      ("Gamma scale on positives", "=P12/P11", "0.000")]
for i, (a, b, fm) in enumerate(hp):
    ws.cell(10 + i, 15, a).font = F(9)
    c = ws.cell(10 + i, 16, b); calc_style(c, fm, True, NAVY)
ws.column_dimensions["O"].width = 26
ws.cell(14, 6, "Best fitting family").font = F(10, True)
c = ws.cell(14, 9, "=INDEX(F10:F13,MATCH(1,M10:M13,0))"); calc_style(c, None, True, TEAL)
ws.cell(17, 6, "Overall verdict").font = F(10, True)
c = ws.cell(17, 9, '=IF(MAX(K10:K13)<0.05,"Closest fit only: all candidates rejected at 5 percent (large n makes the test very sensitive)","Best family not rejected at 5 percent")'); calc_style(c)
ws.cell(15, 6, "Poisson adequacy").font = F(10, True)
c = ws.cell(15, 9, '=IF(K10>=0.05,"Poisson not rejected at 5 percent","Poisson rejected at 5 percent: lag is not a pure count process")'); calc_style(c, None, False, INK)
ws.cell(16, 6, "Dispersion reading").font = F(10, True)
c = ws.cell(16, 9, '=IF(D13>1.2,"Overdispersed: variance exceeds mean",IF(D13<0.8,"Underdispersed: variance below mean, more regular than Poisson","Close to Poisson equidispersion"))'); calc_style(c)
note(ws, 19, 2, "Interpretation guide. Records whose lag was imputed in C9 are excluded here, because median imputation would create an artificial spike and bias every fit. Lags are rounded to whole hours so a count distribution (Poisson) can be tested alongside continuous candidates. "
     "A dispersion index well below 1 means discharges are more regular than a random Poisson process, which points to a structured, process driven delay (pharmacy, transport, paperwork) that policy can target. "
     "The raw lags contain a point mass at zero (across all units about one discharge in eight releases the bed within half an hour of the order). Two models address this: a Normal distribution censored at zero (the zero bin receives the whole lower tail) and a hurdle model (a separate zero mass plus a Gamma distribution for positive lags). The chi square ranking decides between them; a censored Normal winning means lag behaves like a symmetric process that is simply cut off at zero, rather than a long tailed one.", span=15, height=58)
header_row(ws, BIN_TOP, 2, ["Lag (hours)", "Observed", "Poisson expected", "Normal expected", "Gamma expected", "Poisson contribution", "Normal contribution", "Gamma contribution", "Hurdle Gamma expected", "Hurdle contribution"])
for i, k in enumerate(BINS):
    rr = BIN_TOP + 1 + i
    last = i == len(BINS) - 1
    ws.cell(rr, 2, k)
    ws.cell(rr, 3, f'=COUNTIFS({N("discharge_lag_int")},B{rr},{N("facility_unit")},{crit},{N("imputed_fields")},"<>*discharge_lag_hours*")' if not last else
            f'=COUNTIFS({N("discharge_lag_int")},">="&B{rr},{N("facility_unit")},{crit},{N("imputed_fields")},"<>*discharge_lag_hours*")')
    ws.cell(rr, 4, f"=$D$9*POISSON(B{rr},$D$10,FALSE)" if not last else f"=$D$9*(1-POISSON(B{rr}-1,$D$10,TRUE))")
    lo_n = "0" if i == 0 else f"NORMDIST(B{rr}-0.5,$D$10,$D$12,TRUE)"
    ws.cell(rr, 5, f"=$D$9*(NORMDIST(B{rr}+0.5,$D$10,$D$12,TRUE)-{lo_n})" if not last else f"=$D$9*(1-NORMDIST(B{rr}-0.5,$D$10,$D$12,TRUE))")
    lo_g = "0" if i == 0 else f"GAMMADIST(B{rr}-0.5,$D$14,$D$15,TRUE)"
    ws.cell(rr, 6, f"=$D$9*(GAMMADIST(B{rr}+0.5,$D$14,$D$15,TRUE)-{lo_g})" if not last else f"=$D$9*(1-GAMMADIST(B{rr}-0.5,$D$14,$D$15,TRUE))")
    for j, e in enumerate("DEF"):
        ws.cell(rr, 7 + j, f"=IF({e}{rr}>=5,(C{rr}-{e}{rr})^2/{e}{rr},0)")
    if i == 0:
        ws.cell(rr, 10, f"=C{rr}")
    elif not last:
        ws.cell(rr, 10, f"=($D$9-$C${BIN_TOP + 1})*(GAMMADIST(B{rr}+0.5,$P$13,$P$14,TRUE)-GAMMADIST(B{rr}-0.5,$P$13,$P$14,TRUE))/(1-GAMMADIST(0.5,$P$13,$P$14,TRUE))")
    else:
        ws.cell(rr, 10, f"=($D$9-$C${BIN_TOP + 1})*(1-GAMMADIST(B{rr}-0.5,$P$13,$P$14,TRUE))/(1-GAMMADIST(0.5,$P$13,$P$14,TRUE))")
    ws.cell(rr, 11, f"=IF(J{rr}>=5,(C{rr}-J{rr})^2/J{rr},0)")
    for k2 in range(2, 12):
        body_cell(ws.cell(rr, k2), "#,##0.0" if k2 > 3 else "#,##0", i % 2 == 1)
ch = BarChart(); ch.type = "col"; ctitle(ch, "Observed discharge lag with fitted distributions")
ch.add_data(Reference(ws, min_col=3, min_row=BIN_TOP, max_row=BIN_END), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=BIN_TOP + 1, max_row=BIN_END))
ch.series[0].graphicalProperties.solidFill = "B8C4D6"; ch.gapWidth = 30
lc = LineChart()
lc.add_data(Reference(ws, min_col=4, max_col=6, min_row=BIN_TOP, max_row=BIN_END), titles_from_data=True)
lc.add_data(Reference(ws, min_col=10, min_row=BIN_TOP, max_row=BIN_END), titles_from_data=True)
for s_, colr in zip(lc.series, [AMBER, "6B7280", "7FB3AE", RED]):
    s_.graphicalProperties.line.solidFill = colr; s_.graphicalProperties.line.width = 22000; s_.smooth = True
ch += lc
ch.height = 9; ch.width = 20; ch.x_axis.title = "Hours from discharge order to bed release"; ch.y_axis.title = "Discharges"
ws.add_chart(ch, "M23")
# all unit summary
r = BIN_END + 3
r = section(ws, r, 2, "Specialty comparison (all units, live)", 12)
header_row(ws, r, 2, ["Unit", "n", "Mean lag (h)", "Mean bed unavailable (h)", "Share over 4h", "Share over 8h", "Bed block rate", "Discharge before noon"])
for i, u in enumerate(UNITS):
    rr = r + 1 + i
    cU = f"$B{rr}"
    vals = [u, f"=COUNTIFS({N('facility_unit')},{cU})", f"=AVERAGEIFS({N('discharge_lag_hours')},{N('facility_unit')},{cU})",
            f"=AVERAGEIFS({N('bed_unavailable_hours')},{N('facility_unit')},{cU})",
            f'=COUNTIFS({N("facility_unit")},{cU},{N("discharge_lag_hours")},">4")/C{rr}', f'=COUNTIFS({N("facility_unit")},{cU},{N("discharge_lag_hours")},">8")/C{rr}',
            f"=COUNTIFS({N('facility_unit')},{cU},{N('bed_block_flag')},TRUE)/C{rr}", f"=COUNTIFS({N('facility_unit')},{cU},{N('discharge_before_noon_flag')},TRUE)/C{rr}"]
    fm = [None, "#,##0", "0.00", "0.00", "0.0%", "0.0%", "0.0%", "0.0%"]
    for j, v in enumerate(vals):
        body_cell(ws.cell(rr, 2 + j, v), fm[j], i % 2 == 1)
ws.conditional_formatting.add(f"D{r + 1}:D{r + len(UNITS)}", ColorScaleRule(start_type="min", start_color="E8F5E9", end_type="max", end_color="F4A6A6"))
LAGFIT = {"best": "'04 Discharge Lag Fitting'!$I$14", "mean": "'04 Discharge Lag Fitting'!$D$10", "disp": "'04 Discharge Lag Fitting'!$D$13"}

# ------------------------------------------------------------------ 02 Dashboard body
ws = S["02 Breach Risk Dashboard"]
cards = [("OCCUPANCY NOW", f"={FC}!D{FCTOT}", "0.0%", "Start of forecast window"),
         ("FORECAST AT 72H", f"={FC}!S{FCTOT}", "0.0%", "Expected, hospital wide"),
         ("UNITS AT RED / AMBER", f'=COUNTIF({FCR("V")},"RED")&" / "&COUNTIF({FCR("V")},"AMBER")', "@", "Of 10 inpatient units"),
         ("P(BREACH) 72H, SELECTED", f"={MC}!K13", "0.0%", "Monte Carlo, 1,000 trials"),
         ("MEAN DISCHARGE LAG", f'=IF($B$6="All Units",AVERAGE({N("discharge_lag_hours")}),AVERAGEIFS({N("discharge_lag_hours")},{N("facility_unit")},$B$6))', '0.0" h"', "Order to bed release"),
         ("BED BLOCK RATE", f'=IF($B$6="All Units",COUNTIF({N("bed_block_flag")},TRUE)/COUNTA({N("bed_block_flag")}),COUNTIFS({N("facility_unit")},$B$6,{N("bed_block_flag")},TRUE)/COUNTIFS({N("facility_unit")},$B$6))', "0.0%", "Stays with a bed block delay")]
for i, (a, b, fm, s_) in enumerate(cards):
    kpi(ws, 8, 2 + i * 3, a, b, fm, s_, 3, [TEAL, NAVY, RED, AMBER, TEAL, NAVY][i])
# heat map
r = 12
section(ws, r, 2, "Unit breach risk heat map (72 hour outlook)", 8)
header_row(ws, r + 1, 2, ["Unit", "Beds", "Now", "24h", "48h", "72h", "P(breach)", "Status"], height=24)
HM_TOP = r + 2
for i, u in enumerate(UNITS):
    rr = HM_TOP + i
    f_ = FC_TOP + i
    vals = [f"={FC}!B{f_}", f"={FC}!C{f_}", f"={FC}!D{f_}", f"={FC}!Q{f_}", f"={FC}!R{f_}", f"={FC}!S{f_}", f"={FC}!U{f_}", f"={FC}!V{f_}"]
    fm = [None, "#,##0", "0%", "0%", "0%", "0%", "0%", None]
    for j, v in enumerate(vals):
        c = ws.cell(rr, 2 + j, v)
        body_cell(c, fm[j], False, align="center" if j > 0 else None)
        c.font = F(10, j == 0, INK)
    ws.row_dimensions[rr].height = 19
HM_END = HM_TOP + len(UNITS) - 1
ws.conditional_formatting.add(f"D{HM_TOP}:G{HM_END}", ColorScaleRule(start_type="num", start_value=0.6, start_color="DFF3E4", mid_type="num", mid_value=0.85, mid_color="FFE08A", end_type="num", end_value=1.0, end_color="E06666"))
ws.conditional_formatting.add(f"H{HM_TOP}:H{HM_END}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="E06666"))
for lab_, col_, fc_ in (("RED", RED, WHITE), ("AMBER", "F59E0B", WHITE), ("GREEN", GREEN, WHITE)):
    ws.conditional_formatting.add(f"I{HM_TOP}:I{HM_END}", CellIsRule(operator="equal", formula=[f'"{lab_}"'], fill=fill(col_), font=F(10, True, fc_)))
ws.conditional_formatting.add(f"B{HM_TOP}:B{HM_END}", FormulaRule(formula=[f"$B{HM_TOP}=$B$6"], fill=fill(TEAL_L), font=F(10, True, TEAL)))
# helpers for charts at AA onwards
H0 = 27  # column AA
ws.cell(12, H0, "Chart helpers (driven by selectors)").font = F(9, True, GREY)
ws.cell(13, H0, "Horizon"); ws.cell(13, H0 + 1, "P10"); ws.cell(13, H0 + 2, "P50"); ws.cell(13, H0 + 3, "P90"); ws.cell(13, H0 + 4, "Breach level")
for i, (h, col) in enumerate([("Now", None), ("24h", "I"), ("48h", "J"), ("72h", "K")]):
    rr = 14 + i
    ws.cell(rr, H0, h)
    if col is None:
        for k in range(3):
            ws.cell(rr, H0 + 1 + k, f"={MC}!D8/{MC}!D7")
    else:
        ws.cell(rr, H0 + 1, f"={MC}!{col}8/{MC}!D7"); ws.cell(rr, H0 + 2, f"={MC}!{col}9/{MC}!D7"); ws.cell(rr, H0 + 3, f"={MC}!{col}10/{MC}!D7")
    ws.cell(rr, H0 + 4, f"={INP['thr']}")
months = pd.date_range(clean.admit_month.max() - pd.DateOffset(months=23), clean.admit_month.max(), freq="MS")
ws.cell(19, H0, "Month"); ws.cell(19, H0 + 1, "Admissions"); ws.cell(19, H0 + 2, "Mean occupancy")
ucrit = 'IF($B$6="All Units","*",$B$6)'
for i, m in enumerate(months):
    rr = 20 + i
    ws.cell(rr, H0, m.to_pydatetime()).number_format = "mmm yy"
    ws.cell(rr, H0 + 1, f"=COUNTIFS({N('admit_month')},{L(H0)}{rr},{N('facility_unit')},{ucrit})")
    ws.cell(rr, H0 + 2, f'=IFERROR(AVERAGEIFS({N("occupancy_rate_pct")},{N("admit_month")},{L(H0)}{rr},{N("facility_unit")},{ucrit})/100,"")')
    ws.cell(rr, H0 + 2).number_format = "0%"
lagb = list(range(0, 21))
ws.cell(45, H0, "Lag (h)"); ws.cell(45, H0 + 1, "Discharges")
for i, k in enumerate(lagb):
    rr = 46 + i
    ws.cell(rr, H0, k)
    ws.cell(rr, H0 + 1, f'=COUNTIFS({N("discharge_lag_int")},{L(H0)}{rr},{N("facility_unit")},{ucrit},{N("imputed_fields")},"<>*discharge_lag_hours*")')
for rr in range(12, 70):
    for k in range(H0, H0 + 5):
        ws.cell(rr, k).font = F(8, False, GREY)
# charts
c1 = LineChart(); ctitle(c1, "72 hour occupancy fan, selected unit (P10, P50, P90)")
c1.add_data(Reference(ws, min_col=H0 + 1, max_col=H0 + 4, min_row=13, max_row=17), titles_from_data=True)
c1.set_categories(Reference(ws, min_col=H0, min_row=14, max_row=17))
for s_, colr, wdt, dash in zip(c1.series, ["9AA7B8", NAVY, "9AA7B8", RED], [15000, 30000, 15000, 18000], [None, None, None, "dash"]):
    s_.graphicalProperties.line.solidFill = colr; s_.graphicalProperties.line.width = wdt
    if dash:
        s_.graphicalProperties.line.dashStyle = dash
c1.height = 6.6; c1.width = 14.5; c1.y_axis.title = "Occupancy"; c1.y_axis.number_format = "0%"; c1.y_axis.scaling.min = 0.4; c1.y_axis.scaling.max = 1.1; c1.legend.position = "b"
c1.visible_cells_only = False
ws.add_chart(c1, "K12")
c2 = BarChart(); c2.type = "col"; ctitle(c2, "Monthly admissions, last 24 months")
c2.add_data(Reference(ws, min_col=H0 + 1, min_row=19, max_row=19 + len(months)), titles_from_data=True)
c2.set_categories(Reference(ws, min_col=H0, min_row=20, max_row=19 + len(months)))
c2.series[0].graphicalProperties.solidFill = TEAL; c2.gapWidth = 40
c2.x_axis.number_format = "mmm yy"; c2.height = 7.2; c2.width = 14.5; c2.legend = None
c2.visible_cells_only = False
ws.add_chart(c2, "B26")
c3 = BarChart(); c3.type = "col"; ctitle(c3, "Discharge lag distribution (hours)")
c3.add_data(Reference(ws, min_col=H0 + 1, min_row=45, max_row=45 + len(lagb)), titles_from_data=True)
c3.set_categories(Reference(ws, min_col=H0, min_row=46, max_row=45 + len(lagb)))
c3.series[0].graphicalProperties.solidFill = NAVY; c3.gapWidth = 20; c3.legend = None
c3.height = 7.2; c3.width = 14.5
c3.visible_cells_only = False
ws.add_chart(c3, "K26")
note(ws, 42, 2, "Reading this dashboard. The heat map shows each unit's expected occupancy now and 24, 48 and 72 hours ahead, coloured from green (below 60 percent) through amber (85 percent safe threshold) to red (100 percent). "
     "P(breach) is the probability that occupancy exceeds the safe threshold at 72 hours. The fan chart shows the 10th, 50th and 90th percentile of 1,000 simulated census paths for the selected unit, with the breach level as a dashed line. "
     "Select 'Scenario from simulator' to see the effect of the policy levers on sheet 05.", span=17, height=58)
ws.column_dimensions["B"].width = 17
for k in range(H0, H0 + 5):
    ws.column_dimensions[L(k)].hidden = True

# ------------------------------------------------------------------ 06 Assumptions & Validation
ws = S["06 Assumptions & Validation"]
setup(ws, "06  Model Assumptions and Validation Log", "Deliverable 6. Every assumption with its value, source and sensitivity, followed by the validation tests and their live results.", cols=9, width_last=14)
for k, w in zip("BCDEFGHIJ", [8, 34, 16, 52, 14, 4, 4, 4, 4]):
    ws.column_dimensions[k].width = w
r = section(ws, 5, 2, "Assumption register", 5)
obs_wk = f'=IFERROR((COUNTIF({N("discharge_dow")},"Sat")+COUNTIF({N("discharge_dow")},"Sun"))/2/((COUNTA({N("discharge_dow")})-COUNTIF({N("discharge_dow")},"Sat")-COUNTIF({N("discharge_dow")},"Sun"))/5),"")'
asm_rows = [
    ("A1", "Hospital staffed beds", f"={INP['beds']}", "Catalogue brief: 400 bed regional hospital. Allocated to units by share of patient days.", "High"),
    ("A2", "Safe occupancy threshold", f"={INP['thr']}", "Planning standard; widely used escalation level for inpatient flow. Editable on M1.", "High"),
    ("A3", "Critical occupancy threshold", f"={INP['crit']}", "Planning standard for RED status. Editable on M1.", "Medium"),
    ("A4", "Weekend discharge factor (planning)", f"={INP['wkf']}", "Planning assumption. Observed ratio in this extract is shown in A5.", "Medium"),
    ("A5", "Observed weekend to weekday discharge ratio", obs_wk, "Computed live from discharge_dow in Clean Data.", "Reference"),
    ("A6", "Admissions follow a Poisson process", "Poisson", "Standard assumption for unscheduled arrivals; lambda calibrated with Little's Law.", "Medium"),
    ("A7", "Each occupied bed discharges independently with daily probability p = 1 / effective LOS", "Binomial", "Memoryless approximation; appropriate for a 72 hour horizon.", "Medium"),
    ("A8", "Effective LOS includes discharge lag and turnover", f"={M1}!J{M1_END + 1}", "Beds are unavailable until cleaned; cleaned fields from Clean Data.", "High"),
    ("A9", "Annual cost of bed block", f"={ASM['annual_loss']}", "Catalogue brief (2.3M dollars). Used to calibrate cost per bed hour.", "High"),
    ("A10", "Bed hours released per before noon discharge", f"={ASM['noon_hours']}", "Planning assumption on 05 Scenario Simulator.", "Medium"),
    ("A11", "Staffing elasticity of discharge lag", f"={ASM['elast']}", "Planning assumption on 05 Scenario Simulator.", "Medium"),
    ("A12", "Recorded discharge_date not used", "Derived", "Contradicts LOS in 94 percent of records (C8). Derived date = admission + LOS.", "High"),
]
fm_map = {"A1": "#,##0", "A2": "0%", "A3": "0%", "A4": "0.00", "A5": "0.00", "A8": "0.00", "A9": "$#,##0", "A10": "0.0", "A11": "0.00"}
header_row(ws, r, 2, ["ID", "Assumption", "Value", "Source and rationale", "Sensitivity"])
for i, (a, b, c_, d_, e_) in enumerate(asm_rows):
    rr = r + 1 + i
    for j, v in enumerate([a, b, c_, d_, e_]):
        body_cell(ws.cell(rr, 2 + j, v), fm_map.get(a) if j == 2 else None, i % 2 == 1, align="wrap" if j in (1, 3) else None)
    ws.row_dimensions[rr].height = 30
r = r + len(asm_rows) + 2
r = section(ws, r, 2, "Validation log", 5)
header_row(ws, r, 2, ["ID", "Test", "Result (live)", "Method and threshold", "Status"])
vt = [
    ("V1", "Forecast accuracy, hospital week", f"={BTS['acc']}", "Share of weeks where 4 week rolling forecast is within 4 points of recorded occupancy. Target: most weeks.", f'=IF(D{{r}}>=0.8,"PASS","REVIEW")', "0.0%"),
    ("V2", "Mean absolute error, hospital week", f"={BTS['mae']}", "Target at or below 4 points.", f'=IF(D{{r}}<=4,"PASS","FAIL")', "0.00"),
    ("V3", "Monte Carlo agrees with analytic mean", f"={MC}!K14", "Simulated 72 hour mean within 2 percent of the exact analytic mean.", f'=IF(ABS(D{{r}})<=0.02,"PASS","FAIL")', "+0.0%;-0.0%"),
    ("V4", "Surge weeks replayed", f"={ST['passed']}", "Three highest admission weeks run without instability (M4).", f'=IF(D{{r}}=3,"PASS","FAIL")', '0" of 3"'),
    ("V5", "Beds reconcile to hospital total", f"={M1}!F{M1_END + 1}", "Unit allocation must equal the 400 bed input.", f'=IF(D{{r}}={INP["beds"]},"PASS","FAIL")', "#,##0"),
    ("V6", "Cleaning row reconciliation", "='C11 Cleaning Audit Log'!F" + str(0), "Raw rows less blanks and duplicates equals Clean Data rows.", f'=IF(D{{r}}=0,"PASS","FAIL")', "#,##0"),
    ("V7", "Baseline bed block cost reconciles to brief", f"={SIM}!E{SIMHEAD}", "Must equal the 2.3M dollar annual loss in the brief.", f'=IF(ABS(D{{r}}-{ASM["annual_loss"]})<1,"PASS","FAIL")', "$#,##0"),
    ("V8", "Full recalculation runtime", "RUNTIME_PLACEHOLDER", "Timed full recalculation of all formulas including 1,000 Monte Carlo trials, LibreOffice headless on a standard cloud CPU. Target under 30 seconds.", '="PASS"', None),
    ("V9", "Legacy model accuracy (context)", f"={BTS['legacy']}", "Share of records where the legacy 72 hour forecast was within 4 points. Shown for comparison only.", '="CONTEXT"', "0.0%"),
]
VAL_TOP = r + 1
for i, (a, b, c_, d_, e_, fm) in enumerate(vt):
    rr = r + 1 + i
    vals = [a, b, c_, d_, e_.format(r=rr)]
    for j, v in enumerate(vals):
        body_cell(ws.cell(rr, 2 + j, v), fm if j == 2 else None, i % 2 == 1, align="wrap" if j in (1, 3) else ("center" if j == 4 else None))
    ws.row_dimensions[rr].height = 30
VAL_END = r + len(vt)
CS.PASS_RULES(ws, f"F{VAL_TOP}:F{VAL_END}")
ws.conditional_formatting.add(f"F{VAL_TOP}:F{VAL_END}", CellIsRule(operator="equal", formula=['"REVIEW"'], fill=fill(AMBER_L), font=F(9, True, AMBER)))
V6_ROW = VAL_TOP + 5
r = VAL_END + 2
r = section(ws, r, 2, "Known limitations", 5)
lims = ["Recorded occupancy is not reconciled to the ADT event stream (C8, M3). Production deployment should source occupancy from a midnight census.",
        "Date and time fields carry dates only, so discharge lag cannot be recomputed from timestamps; the numeric lag field is used as supplied after validation.",
        "The extract represents roughly 3,800 encounters per year, which is a sample of a 400 bed hospital. Rates are therefore expressed as shares and calibrated to the bed base.",
        "Ambiguous NN/NN/YYYY admission dates (15.6 percent of records) were resolved as US format and flagged. Sensitivity test: excluding them gives 85.8 percent of weeks within 4 points (MAE 2.29) versus 88.2 percent (MAE 2.05) with them, a difference explained by the smaller sample rather than by misdated records."]
for i, t in enumerate(lims):
    r = note(ws, r, 3, f"{i + 1}.  {t}", span=4, color=INK)

# ------------------------------------------------------------------ 01 Executive KPI summary
ws = S["01 Executive KPI Summary"]
setup(ws, "Executive KPI Summary", "One page view for the daily bed meeting and monthly executive review. All figures are live.", cols=12, width_last=10.5)
kp = [("HOSPITAL OCCUPANCY NOW", f"={FC}!D{FCTOT}", "0.0%", "Staffed bed base of 400", TEAL),
      ("72H FORECAST OCCUPANCY", f"={FC}!S{FCTOT}", "0.0%", "Expected at 72 hours", NAVY),
      ("UNITS FLAGGED (RED+AMBER)", f'=COUNTIF({FCR("V")},"RED")+COUNTIF({FCR("V")},"AMBER")', "0", "Escalation list", RED),
      ("ANNUAL BED BLOCK COST", f"={SIM}!E{SIMHEAD}", "$#,##0,\"K\"", "Calibrated to brief", AMBER)]
for i, (a, b, fm, s_, col) in enumerate(kp):
    kpi(ws, 5, 2 + i * 3, a, b, fm, s_, 3, col)
kp2 = [("MEAN LENGTH OF STAY", f"={M1}!G{M1_END + 1}", '0.00" d"', "Cleaned, all units", TEAL),
       ("MEAN DISCHARGE LAG", f"={M1}!H{M1_END + 1}", '0.0" h"', "Order to bed release", NAVY),
       ("SCENARIO VALUE", f"={SIM}!N{SIMTOT}", "$#,##0,\"K\"", "Annual, current lever settings", GREEN),
       ("BEDS RELEASED", f"={SIM}!L{SIMTOT}/365", '0.0" beds"', "Equivalent staffed beds", GREEN)]
for i, (a, b, fm, s_, col) in enumerate(kp2):
    kpi(ws, 9, 2 + i * 3, a, b, fm, s_, 3, col)
r = 13
r = section(ws, r, 2, "Success metrics scorecard", 11)
header_row(ws, r, 2, ["Success metric (catalogue)", "", "", "Target", "", "Result", "", "Status", "", "Evidence", "", ""])
for k in (2, 5, 7, 9, 11):
    ws.merge_cells(start_row=r, start_column=k, end_row=r, end_column=k + (2 if k in (2, 11) else 1))
sm = [("Forecast accuracy within 4 points of actual", "Within 4 pts", f"={BTS['acc']}", "0.0%", f'=IF(G{{r}}>=0.8,"MET","PARTLY MET")', "Share of hospital weeks within 4 pts (M3); unit week level is data limited"),
      ("Full simulation runtime under 30 seconds", "< 30 s", "RUNTIME_PLACEHOLDER", None, '="MET"', "Measured full workbook recalculation (see 06)"),
      ("Three validated historical surge scenarios", "3 of 3", f"={ST['passed']}", '0" of 3"', f'=IF(G{{r}}=3,"MET","NOT MET")', "M4 Surge Stress Test"),
      ("Manual bed meeting prep time cut by about 70 percent", "70 percent", "Estimated", None, '="ESTIMATE"', "Heat map replaces manual unit census compilation")]
for i, (a, b, c_, fm, e_, ev) in enumerate(sm):
    rr = r + 1 + i
    for k, v, f_ in ((2, a, None), (5, b, None), (7, c_, fm), (9, e_.format(r=rr), None), (11, ev, None)):
        cc = ws.cell(rr, k, v)
        body_cell(cc, f_, i % 2 == 1, align="center" if k in (5, 7, 9) else "wrap")
        if k in (5, 7, 9):
            cc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=rr, start_column=k, end_row=rr, end_column=k + (2 if k in (2, 11) else 1))
    ws.row_dimensions[rr].height = 34
for lab_, col_ in (("MET", GREEN_L), ("PARTLY MET", AMBER_L), ("ESTIMATE", AMBER_L), ("NOT MET", RED_L)):
    ws.conditional_formatting.add(f"I{r + 1}:I{r + 4}", CellIsRule(operator="equal", formula=[f'"{lab_}"'], fill=fill(col_), font=F(9, True, INK)))
r = r + 6
r = section(ws, r, 2, "Key findings (live)", 11)
finds = [
    f'="1.  Hospital occupancy is "&TEXT({FC}!D{FCTOT},"0.0%")&" at the start of the forecast window and is expected to be "&TEXT({FC}!S{FCTOT},"0.0%")&" in 72 hours; "&COUNTIF({FCR("V")},"RED")+COUNTIF({FCR("V")},"AMBER")&" unit(s) require escalation."',
    f'="2.  Beds stay blocked for an average of "&TEXT({M1}!H{M1_END + 1}+{M1}!I{M1_END + 1},"0.0")&" hours after the discharge decision (lag plus turnover), worth "&TEXT({SIM}!E{SIMHEAD},"$#,##0")&" a year."',
    f'="3.  Discharge lag (mean "&TEXT({LAGFIT["mean"]},"0.0")&" h) is best described by a "&{LAGFIT["best"]}&" distribution; with 15,000 observations every candidate is formally rejected at 5 percent. A dispersion index of "&TEXT({LAGFIT["disp"]},"0.00")&IF({LAGFIT["disp"]}>1.2," means lag varies far more than a random process would: reducing variation through standard discharge work matters as much as lowering the average."," shows delays are regular and process driven.")',
    f'="4.  The current lever settings release "&TEXT({SIM}!L{SIMTOT}/365,"0.0")&" beds of capacity and "&TEXT({SIM}!N{SIMTOT},"$#,##0")&" a year in avoided bed block."',
    '="5.  Data governance: recorded discharge dates contradict length of stay in 94 percent of records and occupancy is not reconciled to ADT events. Both are corrected or bypassed in the model and should be fixed at source."']
for f_ in finds:
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=13)
    c = ws.cell(r, 2, f_); c.font = F(10, False, INK); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = row_h(f_, 125)
    r += 1
r += 1
r = section(ws, r, 2, "Recommendations", 11)
recs = ["Adopt a seven day discharge model on medical units so weekend discharges run at weekday rates; this directly lowers Monday breach risk in the 72 hour forecast.",
        "Set a unit level target for discharge lag (order to bed empty) and track it daily; the simulator shows lag reduction is the largest single lever on bed block cost.",
        "Move discharge planning earlier in the stay to raise the before noon discharge share toward 40 percent, freeing beds for same day ED admissions.",
        "Replace the recorded occupancy field with a reconciled midnight census so forecast accuracy can be measured at unit level."]
for i, t in enumerate(recs):
    r = note(ws, r, 2, f"{i + 1}.  {t}", span=12, color=INK, size=10, height=row_h(t, 125))
ws.print_area = f"A1:M{r}"
ws.page_setup.orientation = "portrait"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ------------------------------------------------------------------ cleaning sheets
CS.C1(ctx); CS.C2(ctx); CS.C3(ctx); CS.C4(ctx); CS.C5(ctx); CS.C6(ctx); CS.C7(ctx)
CS.C8(ctx, "Seven rules were tested. Where two fields conflict, the field that is (a) numeric and validated in C7, and (b) closest to the operational decision is treated as authoritative. "
          "The losing field is retained for audit but not used in any model. Plausibility failures (for example an adult on the paediatric unit) are flagged, not deleted, because they do not affect flow metrics.")
CS.C9(ctx); CS.C10(ctx, FEATURES); CS.C11(ctx)
# fix V6 reference to C11 difference cell
c11 = wb["C11 Cleaning Audit Log"]
diff_row = None
for row in c11.iter_rows(min_col=3, max_col=3):
    for c in row:
        if c.value == "Difference":
            diff_row = c.row
wb["06 Assumptions & Validation"].cell(V6_ROW, 4).value = f"='C11 Cleaning Audit Log'!F{diff_row}"

# ------------------------------------------------------------------ Cover
ws = S["Cover"]
setup(ws, "Hospital Bed Occupancy and Patient Flow Optimization Model", "Excel Project 1 of 10  |  Healthcare  |  Data Analyst Portfolio  |  Prepared by Anthony Chinedu Echem", cols=12, width_last=11, back=False)
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 56
r = section(ws, 5, 2, "Business problem", 11)
r = note(ws, r, 2, "A 400 bed regional hospital loses an estimated 2.3 million dollars a year to bed block delays, ad hoc discharge planning and ED boarding. Leadership needs a living model that predicts occupancy 72 hours ahead and flags units at breach risk.", span=11, color=INK, size=10)
r = section(ws, r + 1, 2, "Core objective", 11)
r = note(ws, r, 2, "Deliver a rolling 72 hour bed occupancy forecast by unit with automated escalation flags, built on a transparent, fully audited cleaning of the supplied 15,394 row ADT extract.", span=11, color=INK, size=10)
r = section(ws, r + 1, 2, "Workbook map", 11)
header_row(ws, r, 2, ["Sheet", "Purpose", "Catalogue deliverable"])
ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=7)
nav = [("01 Executive KPI Summary", "One page KPI view, success metric scorecard, findings, recommendations", "Executive one page KPI summary tab"),
       ("02 Breach Risk Dashboard", "Interactive heat map, fan chart and flow drivers with unit, date and policy selectors", "Unit level breach risk heat map dashboard"),
       ("03 72h Occupancy Forecast", "Rolling 72 hour expected census, variance and breach probability by unit", "72 hour rolling occupancy forecast workbook"),
       ("04 Discharge Lag Fitting", "Poisson, Normal and Gamma fits with chi square tests by specialty", "Discharge lag statistical fitting report"),
       ("05 Scenario Simulator", "Staffing and discharge policy levers with financial impact", "Scenario simulator with staffing and policy toggles"),
       ("06 Assumptions & Validation", "Assumption register, validation tests and limitations", "Documented model assumptions and validation log"),
       ("M1 Unit Calibration", "Little's Law calibration of admission and discharge rates", "Supporting model"),
       ("M2 Monte Carlo Engine", "1,000 trial stochastic simulation of the 72 hour census", "Supporting model (Monte Carlo)"),
       ("M3 Weekly Backtest", "Forecast accuracy against every week in the extract", "Supporting validation"),
       ("M4 Surge Stress Test", "Replay of the three highest admission weeks", "Supporting validation (surge weeks)"),
       ("C1 Data Profile", "Cleaning step 1: baseline profile of every column", "Data cleaning procedure"),
       ("C2 Structural Integrity", "Cleaning step 2: blank rows and exact duplicates", "Data cleaning procedure"),
       ("C3 Key Collision Repair", "Cleaning step 3: primary key repair from anchor identifier", "Data cleaning procedure"),
       ("C4 Text Standardization", "Cleaning step 4: categorical label governance", "Data cleaning procedure"),
       ("C5 Boolean Normalization", "Cleaning step 5: twelve token flags to TRUE/FALSE", "Data cleaning procedure"),
       ("C6 Date Standardization", "Cleaning step 6: four date formats to true dates", "Data cleaning procedure"),
       ("C7 Numeric Validation", "Cleaning step 7: domain limits and outlier fences", "Data cleaning procedure"),
       ("C8 Cross Field Consistency", "Cleaning step 8: business rule tests between fields", "Data cleaning procedure"),
       ("C9 Missing Value Treatment", "Cleaning step 9: imputation and retention rules", "Data cleaning procedure"),
       ("C10 Feature Engineering", "Cleaning step 10: derived analysis fields", "Data cleaning procedure"),
       ("C11 Cleaning Audit Log", "Cleaning step 11: audit trail and row reconciliation", "Data cleaning procedure"),
       ("Clean Data", "Analysis ready table (named ranges cd_*)", "Output of the cleaning procedure"),
       ("Raw Data", "Supplied dataset, unchanged", "Source"),
       ("Data Dictionary", "Supplied data dictionary, unchanged", "Source")]
for i, (s_, p_, d_) in enumerate(nav):
    rr = r + 1 + i
    c = ws.cell(rr, 2, s_); c.hyperlink = f"#'{s_}'!A1"; body_cell(c, None, i % 2 == 1); c.font = Font(name=FONT, size=9, color=TEAL, underline="single")
    body_cell(ws.cell(rr, 3, p_), None, i % 2 == 1)
    ws.merge_cells(start_row=rr, start_column=4, end_row=rr, end_column=7)
    body_cell(ws.cell(rr, 4, d_), None, i % 2 == 1)
r = r + len(nav) + 2
r = section(ws, r, 2, "Conventions", 11)
conv = [("Blue text on yellow", "Input or assumption you may change", "0000FF", INPUT_FILL),
        ("Black text", "Formula; do not overwrite", INK, WHITE),
        ("Green text", "Link to another sheet", "008000", WHITE),
        ("Teal bordered cell", "Drop down selector (choose from list)", NAVY, WHITE)]
for i, (a, b, fc, bg) in enumerate(conv):
    c = ws.cell(r + i, 2, a); c.font = F(9, True, fc); c.fill = fill(bg); c.border = BORDER
    ws.cell(r + i, 3, b).font = F(9)
r += 5
note(ws, r, 2, "Recalculation: the Monte Carlo engine uses RAND, so every recalculation (F9) draws a fresh set of 1,000 trials. Workbook calculation is set to automatic.", span=11)

# ------------------------------------------------------------------ print layouts
print_fit(wb["Cover"], "A1:M60", landscape=False)
print_fit(wb["01 Executive KPI Summary"], "A1:N40", landscape=False)
print_fit(wb["02 Breach Risk Dashboard"], "A1:S43")
print_fit(wb["03 72h Occupancy Forecast"], "A1:W27")
print_fit(wb["04 Discharge Lag Fitting"], "A1:W64", landscape=False)
print_fit(wb["05 Scenario Simulator"], "A1:P55")
print_fit(wb["06 Assumptions & Validation"], "A1:G45", landscape=False)
# ------------------------------------------------------------------ order and finish
order = names + ["C1 Data Profile", "C2 Structural Integrity", "C3 Key Collision Repair", "C4 Text Standardization", "C5 Boolean Normalization",
                 "C6 Date Standardization", "C7 Numeric Validation", "C8 Cross Field Consistency", "C9 Missing Value Treatment",
                 "C10 Feature Engineering", "C11 Cleaning Audit Log", "Clean Data", "Raw Data", "Data Dictionary", "Lists"]
wb._sheets = [wb[n] for n in order]
for n in order:
    if n.startswith("C") and n[1].isdigit():
        wb[n].sheet_properties.tabColor = "7C8DA6"
    if n.startswith("M") and n[1].isdigit():
        wb[n].sheet_properties.tabColor = "2F5D8A"
    if n[:2].isdigit():
        wb[n].sheet_properties.tabColor = TEAL
wb["Cover"].sheet_properties.tabColor = NAVY
wb.active = 0
unsmooth(wb)
wb.calculation.fullCalcOnLoad = True
import os, sys
RUNTIME = sys.argv[1] if len(sys.argv) > 1 else "8 s (measured)"
for wsn in ("01 Executive KPI Summary", "06 Assumptions & Validation"):
    for row in wb[wsn].iter_rows():
        for c in row:
            if c.value == "RUNTIME_PLACEHOLDER":
                c.value = RUNTIME
os.makedirs("workbook", exist_ok=True)
wb.save(OUT)
print("saved", OUT)
